#!/usr/bin/env python3
"""Regenerate the .xlsx fixtures that check_tables.py reads.

    python3 test/python/fixtures/make_table_workbooks.py

The only non-stdlib thing in the repo, and deliberately outside the test run: the
workbooks are committed, and this script exists so a defect planted in one can be
changed without hand-editing a zip. It needs ``openpyxl``; the test suite does not.

openpyxl is used rather than a hand-rolled writer on purpose — the reader in
check_tables.py has to cope with the XML a real spreadsheet library emits (shared
strings, relationship ids, a formula with no cached value), not with XML written to
suit it.
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook

HERE = Path(__file__).resolve().parent
GOOD = HERE / "good" / "figs" / "ColMemo_Tables.xlsx"
BAD = HERE / "bad" / "figs" / "Bad_Tables.xlsx"
EXTRA = HERE / "bad" / "figs" / "Extra_Tables.xlsx"

CAPTION_1 = "Parameters of the memory model, with the range each was swept over."
CAPTION_S1 = "Climate information recovered from each documentary archive."


def sheet(wb, title, rows):
    ws = wb.create_sheet(title)
    for row in rows:
        ws.append(row)
    return ws


def build_good() -> None:
    wb = Workbook()
    wb.remove(wb.active)
    sheet(wb, "_index", [
        ["sheet", "label", "caption", "source", "updated"],
        ["Table 1", "tbl:params", CAPTION_1, "results.json:sweep_bounds", "2026-08-12"],
        ["Table S1", "tbl:archives", CAPTION_S1, "reports/build_archives.py", "2026-08-12"],
    ])
    sheet(wb, "Table 1", [
        [CAPTION_1],
        ["Parameter", "Symbol", "Default", "Swept range"],
        ["Memory length", "$L_{\\max}$", 20, "5-50"],
        ["Loss rate", "$\\lambda$", 0.05, "0.01-0.2"],
        ["Climate weight", "$\\phi$", 0.4, "0.1-0.9"],
        ["New arrivals", "$N_{\\text{new}}$", 12, "0-40"],
    ])
    sheet(wb, "Table S1", [
        [CAPTION_S1],
        ["Archive", "Region", "Span (CE)", "Records (n)"],
        ["Local gazetteers", "North China", "1470-1949", 208],
        ["Court diaries", "Capital", "1644-1911", 96],
        ["Tree rings", "Northwest", "1300-2000", 41],
        ["Ice cores", "Tibet", "800-1990", 12],
    ])
    wb.save(GOOD)


def build_bad() -> None:
    wb = Workbook()
    wb.remove(wb.active)
    # No _index sheet at all -- the planted sheet-index defect.
    ws = sheet(wb, "Table 1", [
        [CAPTION_1],
        ["Parameter", "Default", "Swept range"],
        ["Memory length", 20, "5-50"],
        ["Loss rate", None, "0.01-0.2"],
    ])
    ws["B4"] = "=0.01*5"          # a formula: openpyxl stores no cached value
    sheet(wb, "Table 2", [
        ["Model fit against the two withheld archives."],
        ["Archive", "RMSE"],
        ["Tree rings", 0.41],
        ["Ice cores", 0.52],
    ])
    sheet(wb, "Table S3", [
        ["Grading determinations, with the rule applied to each."],
        ["Case", "Grade", "Rule"],
        ["1470 drought", "severe", "below -1.17 sigma"],
        ["1523 flood", "moderate", "above 1.17 sigma"],
        ["1602 drought", "severe", "below -1.17 sigma"],
    ])
    wb.save(BAD)

    # A second workbook on the same resource path -- the workbook-single defect.
    wb2 = Workbook()
    wb2.active.title = "Table 4"
    wb2.active.append(["Left over from an earlier draft."])
    wb2.save(EXTRA)


if __name__ == "__main__":
    for path in (GOOD, BAD, EXTRA):
        path.parent.mkdir(parents=True, exist_ok=True)
    build_good()
    build_bad()
    print(f"wrote {GOOD.name}, {BAD.name}, {EXTRA.name}")
